#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the AMZ Visuals DEVELOPER to-do as an English Excel workbook."""
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ORANGE = "F06524"; DARK = "221F1F"
HEADER_FILL = PatternFill("solid", fgColor=ORANGE)
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
TITLE_FONT = Font(bold=True, color=DARK, size=16)
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")
CENTER = Alignment(horizontal="center", vertical="top")
thin = Side(style="thin", color="DDDDDD")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

PRIO_FILL = {"Blocker": PatternFill("solid", fgColor="E74C3C"),
             "High": PatternFill("solid", fgColor="E67E22"),
             "Medium": PatternFill("solid", fgColor="F1C40F"),
             "Low": PatternFill("solid", fgColor="BDC3C7")}
PRIO_FONT = {"Blocker": Font(bold=True, color="FFFFFF"),
             "High": Font(bold=True, color="FFFFFF"),
             "Medium": Font(bold=True, color=DARK),
             "Low": Font(bold=True, color=DARK)}
TYPE_FILL = {"Bug": PatternFill("solid", fgColor="FADBD8"),
             "Adjustment": PatternFill("solid", fgColor="FDEBD0"),
             "New": PatternFill("solid", fgColor="D6EAF8"),
             "Architecture": PatternFill("solid", fgColor="E8DAEF")}

# ID, Priority, Topic, Type, Task, Information, Source, Sprint
rows = [
 # --- BLOCKERS ---
 ("B1","Blocker","Onboarding","Bug","Signup broken / no confirmation email",
  "New users cannot register; no confirmation email (not in spam). Check Supabase Auth + SMTP/email provider (templates, sender domain, rate limits). Test full signup with a fresh email. Add proper error/success states.","02:24-03:37","1"),
 ("B2","Blocker","Payments","Bug","Checkout does not complete (stays 'pending')",
  "UI shows 'payment successful' but order stays 'payment pending'; even with a 100% code. Verify Stripe webhook (checkout.session.completed / payment_intent.succeeded) is received and sets order status. Handle the 100%-discount (0 EUR) path separately. Fix the long 'finish order' loading.","13:00-15:14","1"),
 ("B3","Blocker","AI Module","Bug","AI image generation not running",
  "AI generation does not start at all. Check AI provider connection/keys/quota; log errors. Restore flow: product -> type 'Main Image' -> reference image -> generate.","23:30","1"),
 # --- HIGH ---
 ("H1","High","Order Flow","Bug","State loss on tab switch",
  "While creating a product/order, switching browser tab (to copy a competitor ASIN) sometimes logs the user out / loses progress. Autosave the product + order draft.","11:32-12:19","2"),
 ("H2","High","Order Status","Bug","Wrong status 'in progress' after approve-anyway",
  "After 'approve anyway', the status must be 'finalizing', not 'in progress'.","17:36","2"),
 ("H3","High","Designer","Bug","Designer dashboard does not auto-update",
  "Designer must reload to see new state. Add live / auto refresh.","18:24","2"),
 ("H4","High","Designer","Adjustment","Approval state misleading for designer",
  "With 'approve anyway' (no per-item approval) the designer sees 'approved' although items are not approved and user comments exist. Make the real review status unambiguous; surface user comments prominently.","19:20","2"),
 ("H5","High","Notifications","New","No completion notification / email",
  "On 'mark complete' the client gets no in-app notification and no email. Notification must also be set to unread; send the client an email.","21:43","2"),
 # --- UX / COPY / VISUAL ---
 ("U1","Medium","Copy","Adjustment","Fix 'few projects' text",
  "After payment it says 'few projects' -> should read 'View project' / 'View visual order'.","14:11","2"),
 ("U2","Medium","Review UI","Adjustment","Order overview has too much copy",
  "Overview is text-heavy; the key info is hard to see. Simplify, sharpen the hierarchy.","15:56","2"),
 ("U3","Medium","Review UI","Adjustment","'Pending review' once per product, not per image",
  "Show 'Pending review' once next to the product name (e.g. 'Main V1 - Pending review'), not on every image.","16:13","2"),
 ("U4","Medium","Review UI","Adjustment","Approved image clearly shows 'Approved'",
  "After approving a single image it must clearly show 'Approved'; the approve button then disappears for it.","17:15","3"),
 ("U5","Medium","Order Flow","Adjustment","Consider removing 'Review order' step",
  "The separate 'Review order' step seems unnecessary -> consider going straight to payment.","12:39","3"),
 ("U6","Medium","Review UI","Adjustment","Remove 'Feedback' wording for client",
  "On a finished project, remove anything labeled 'Feedback' for the client; replace with 'Not happy? Reach out to us' + contact email.","22:07","2"),
 ("U7","Medium","Download","Adjustment","Improve download discoverability",
  "From a finished order the download is hard to find. Add a clear hint/link: 'To download -> Visual Library'.","22:21","3"),
 ("U8","Low","Brand Form","Adjustment","Hint: brand guide replaces colors/fonts",
  "Add a hint in the brand form: 'If a brand guide is uploaded, colors/fonts are optional.'","05:13","5"),
 ("V1","Low","Copy","Bug","Typo 'No Gas' -> 'No-Gos'",
  "Brand form shows 'No Gas'; should be 'No-Gos'.","Frame f9","2"),
 ("V2","Medium","UI/Design","Adjustment","Raise global text contrast",
  "Info banners, helper texts and placeholders are light-gray on light; landing subheadline / CVR lines are barely readable. Raise contrast to WCAG AA.","Frames f1,f7,f9,f19,f46","2"),
 ("V3","Medium","UI/Design","Bug","Chat widget overlaps CTAs / modals",
  "The 'Got any questions?' chat overlaps the primary CTA (e.g. 'Create Brand') and the upload modal. Fix position / z-index.","Frames f7,f9,f52","2"),
 ("V4","Medium","Stability","Bug","Empty / broken render states",
  "Several fully blank app screens (Supabase URLs visible) during the flow. Add loading / error states. Related to B1/B2/B3.","Frames f3,f50,f54","2"),
 # --- AI MODULE ---
 ("A2","High","AI Module","New","AI prompt agent for main image",
  "From product data + uploaded image, auto-build a strong main-image concept: combine a relevant element (e.g. product + chicken), suggest a badge/hook (e.g. '40% Protein'), make food look more appetizing where it makes sense. Target: click-through-strong Amazon main images.","24:39-26:03","3"),
 ("A3","High","AI Module","New","Product fidelity (compositing / inpainting)",
  "Preserve the real product from the photo; only generate scene/background. No pure text-to-image (it would alter the product -> unusable for Amazon).","spec","3"),
 ("A4","Medium","AI Module","New","Text via overlays / templates",
  "Render badge / infographic text as editable overlay layers, not via AI (legibility).","spec","3"),
 ("A5","Medium","AI Module","New","Amazon conformity guardrail",
  "Validate dimensions, white background for hero, no invalid claims/badges before download.","spec","3"),
 ("A6","Medium","AI Module","Architecture","ImageProvider abstraction",
  "Wrap the image model behind an ImageProvider interface; models change often.","spec","3"),
 # --- AD VIDEO ---
 ("AV1","Medium","Ad-Video","New","Ad-video stage 1 (template / motion)",
  "Template/motion-based video (images + animation + text + music, e.g. Remotion) + render worker. Export presets: Amazon Listing / Sponsored Brands / 9:16 / 1:1. Buyable in-tool (subscription/credits).","spec","4"),
 ("AV2","Low","Ad-Video","New","Ad-video stage 2 (generative)",
  "Generative AI video behind a VideoProvider interface (Kling/Runway/Veo) as a premium add-on, later.","spec","later"),
 # --- PARTNER PROGRAM (minimal dev) ---
 ("P1","High","Partner Program","New","Connect Tolt to Stripe",
  "Set up Tolt on Stripe: tracking links, cookie attribution, partner dashboards, 25% lifetime commission, payouts via Wise/Payoneer/PayPal. Tolt reads the real Stripe charges, so lifetime + net-of-discount is handled natively.","decision 13.06","4"),
 ("P2","High","Partner Program","New","Agency attribution at signup",
  "Store ?ref=CODE from the affiliate link in a cookie and attach it to the user at signup (referred_by_agency). Add an optional signup field 'Referred by an agency? (name/code)' as a fallback.","decision 13.06","4"),
 ("P3","Medium","Partner Program","New","Partner dashboard (lean)",
  "Read-only view for agencies: referred clients, revenue, commission earned, payout status. v1 can come from Tolt; native later.","decision 13.06","4"),
 ("P6","Medium","Partner Program","New","Agency-managed mode (optional)",
  "For agencies ordering in their own account for clients: 25% standing discount instead of commission; clients as brands. Uses existing structures.","decision 13.06","later"),
 # --- ARCHITECTURE (cross-cutting) ---
 ("ARCH1","High","Architecture","Architecture","Multi-tenancy / org_id from day 1",
  "Introduce org_id / tenant_id in the data model, auth & storage now (even if full white-label comes later). Retrofitting is expensive.","handoff","1"),
 ("ARCH2","Medium","Architecture","Architecture","Async jobs / queue",
  "Run AI generation & video rendering via a queue + status + webhooks; never synchronously inside the request.","handoff","3"),
 ("ARCH3","Low","Architecture","Architecture","Logging / monitoring",
  "Ensure logging/monitoring so failures are visible ('a lot of things have been downed').","-","1"),
]

wb = Workbook()
ws = wb.active
ws.title = "Developer To-Do"
headers = ["ID","Priority","Topic","Type","Task","Information","Source","Sprint","Effort","Owner","Status"]

ws.merge_cells("A1:K1")
ws["A1"] = "AMZ Visuals - Developer To-Do (Relaunch)"
ws["A1"].font = TITLE_FONT; ws.row_dimensions[1].height = 26
ws.merge_cells("A2:K2")
ws["A2"] = ("Live tool >> old prototype. Backend = Supabase, /designer route exists. "
            "Relaunch = fix + finish. Work top-down. Source = Loom timestamp / video frame. "
            "Fill Effort/Owner/Status yourself.")
ws["A2"].font = Font(italic=True, color="555555", size=9)
ws["A2"].alignment = WRAP; ws.row_dimensions[2].height = 26

hr = 3
for c,h in enumerate(headers, start=1):
    cell = ws.cell(row=hr, column=c, value=h)
    cell.fill = HEADER_FILL; cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER
ws.row_dimensions[hr].height = 22

r = hr + 1
for (rid,prio,topic,typ,task,info,src,sprint) in rows:
    ws.cell(row=r, column=1, value=rid).alignment = CENTER
    pc = ws.cell(row=r, column=2, value=prio); pc.alignment = CENTER
    pc.fill = PRIO_FILL[prio]; pc.font = PRIO_FONT[prio]
    ws.cell(row=r, column=3, value=topic).alignment = TOP
    tc = ws.cell(row=r, column=4, value=typ); tc.alignment = CENTER
    if typ in TYPE_FILL: tc.fill = TYPE_FILL[typ]
    tk = ws.cell(row=r, column=5, value=task); tk.alignment = WRAP; tk.font = Font(bold=True, color=DARK)
    ws.cell(row=r, column=6, value=info).alignment = WRAP
    ws.cell(row=r, column=7, value=src).alignment = WRAP
    ws.cell(row=r, column=8, value=sprint).alignment = CENTER
    ws.cell(row=r, column=9, value="").alignment = TOP
    ws.cell(row=r, column=10, value="").alignment = TOP
    ws.cell(row=r, column=11, value="Open").alignment = CENTER
    for c in range(1,12):
        ws.cell(row=r, column=c).border = BORDER
    r += 1

for i,w in enumerate([7,10,15,12,34,62,16,8,9,14,9], start=1):
    ws.column_dimensions[get_column_letter(i)].width = w
ws.freeze_panes = "A4"
ws.auto_filter.ref = f"A{hr}:K{r-1}"

# Sheet 2: Context & Architecture
ws2 = wb.create_sheet("Context & Architecture")
notes = [
 ("Key fact","The live deployed tool is far more built than the old prototype (sb1-oc1x2hve). Signup, payment, designer dashboard with versioning, visual library and an AI module already exist. The 3 priorities = FIX + FINISH, not greenfield."),
 ("Stack (live)","Backend = Supabase (*.supabase.co URLs visible in frames). Own /designer route. For B1/B2 check Supabase Auth/RLS/webhooks/keys and the Stripe webhook first."),
 ("Multi-tenancy","Add org_id/tenant_id from day 1 (data model, auth, storage) - even though full white-label is later. Retrofitting is expensive."),
 ("AI provider abstraction","Wrap image/video models behind ImageProvider / VideoProvider interfaces; models change monthly."),
 ("Async","AI generation & video rendering run via queue + status + webhooks, never synchronously."),
 ("Partner program","Affiliate model (not white-label). Tool = Tolt on Stripe. Commission = 25% lifetime on the actual paid amount. Minimal custom dev = ref-cookie capture + optional signup agency field."),
 ("Suggested sprints","Sprint 1: prerequisites + blockers (B1,B2,B3) + ARCH1/ARCH3 -> relaunch-ready. Sprint 2: high bugs (H1-H5) + main UX/visual (U1,U2,U3,U6,V1,V2,V3,V4). Sprint 3: finalize AI module. Sprint 4: partner program (Tolt) + ad-video stage 1. Remaining UX ongoing."),
 ("Full docs","Repo JonasK01/amzvisualsai, branch claude/hopeful-gauss-41dvp7: docs/01-10. Narrative dev brief = docs/10-developer-todo.md."),
]
ws2.merge_cells("A1:B1"); ws2["A1"] = "Context & Architecture"; ws2["A1"].font = TITLE_FONT
ws2.row_dimensions[1].height = 26
for c,h in enumerate(["Topic","Note"], start=1):
    cell = ws2.cell(row=2, column=c, value=h)
    cell.fill = HEADER_FILL; cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center"); cell.border = BORDER
rr = 3
for (t,n) in notes:
    ws2.cell(row=rr, column=1, value=t).alignment = WRAP
    ws2.cell(row=rr, column=1).font = Font(bold=True, color=DARK)
    ws2.cell(row=rr, column=2, value=n).alignment = WRAP
    for c in range(1,3): ws2.cell(row=rr, column=c).border = BORDER
    rr += 1
ws2.column_dimensions["A"].width = 24; ws2.column_dimensions["B"].width = 100
ws2.freeze_panes = "A3"

out = "/home/user/amzvisualsai/AMZVisuals_Developer_ToDo_EN.xlsx"
wb.save(out)
print("Saved:", out, "| rows:", len(rows))
